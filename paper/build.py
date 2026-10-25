"""Assemble the Intelligent Oncology submission package (Step 5.2). CPU only.

    python paper/build.py            # writes paper/submission/ and compiles the manuscript

Steps
  1. references.tex in order of first citation (AMA, superscript numerals via natbib)
  2. paper/submission/ with main.tex, references.tex, figures/Figure1-6 (PDF and 600 dpi PNG), highlights.txt;
     Figure1 and the graphical abstract are TikZ sources in paper/figures_src, the others come from figures/
  3. manuscript.pdf compiled with the official elsarticle class in paper/elsarticle (needs pdflatex)
  4. manuscript.docx through pandoc, single column, left aligned, double spaced, line numbers
  5. supplement tables (scripts/make_supplement.py) and supplement.pdf
"""
from __future__ import annotations

import os
import re
import shutil
import zipfile
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(HERE, "submission")
sys.path.insert(0, HERE)
from references import REFS  # noqa: E402

FIGS = {"Figure2": "fig1_privacy_utility", "Figure3": "fig4_h5_noise", "Figure4": "fig2_lira_roc",
        "Figure5": "fig3_explanations", "Figure6": "fig6_gradcam_examples"}
TIKZ = {"Figure1": "framework"}
ALG_WORDS = {"Require": "Input.", "Ensure": "Output.", "State": "", "For": "for", "EndFor": "end for",
             "If": "if", "ElsIf": "else if", "Else": "else", "EndIf": "end if"}


def tikz_pdf(src: str, dest_dir: str, name: str) -> None:
    """Compile a standalone TikZ figure to PDF and a 600 dpi PNG."""
    d = os.path.join(HERE, "figures_src")
    run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", f"{src}.tex"], d)
    shutil.copy(os.path.join(d, f"{src}.pdf"), os.path.join(dest_dir, f"{name}.pdf"))
    run(["pdftoppm", "-r", "600", "-png", "-singlefile", os.path.join(d, f"{src}.pdf"),
         os.path.join(dest_dir, name)], d)


def algorithm_text(m) -> str:
    """An algorithm environment as a numbered caption and indented plain lines, for the Word version."""
    body = m.group(0)
    cap = re.search(r"\\caption\{(.*?)\}\\label", body, re.S).group(1)
    lines, depth = [], 0
    body = re.sub(r"\n[ \t]*(?=[^\\ \t\n])", " ", body)      # join wrapped continuation lines
    for raw in re.findall(r"\\(Require|Ensure|State|For|EndFor|If|ElsIf|Else|EndIf)\b(.*)", body):
        cmd, rest = raw
        if cmd in ("EndFor", "EndIf", "ElsIf", "Else"):
            depth -= 1
        rest = rest.strip()
        rest = re.sub(r"^\{(.*)\}$", r"\1", rest) if cmd in ("For", "If", "ElsIf") else rest
        rest = re.sub(r"\\Comment\{(.*?)\}", r" (\1)", rest)
        tail = " do" if cmd == "For" else (" then" if cmd in ("If", "ElsIf") else "")
        text = (ALG_WORDS[cmd] + " " + rest).strip() + tail
        lines.append("~" * (4 * max(depth, 0)) + text)
        if cmd in ("For", "If", "ElsIf", "Else"):
            depth += 1
    items = "\n".join(f"\\item {t}" for t in lines)
    return f"\\paragraph{{ALGCAP {cap}}}\n\\begin{{enumerate}}\n{items}\n\\end{{enumerate}}\n"


def cite_order(tex: str) -> list[str]:
    order = []
    for grp in re.findall(r"\\cite\{([^}]*)\}", tex):
        for k in (x.strip() for x in grp.split(",")):
            if k not in order:
                order.append(k)
    missing = [k for k in order if k not in REFS]
    unused = [k for k in REFS if k not in order]
    if missing or unused:
        sys.exit(f"citation mismatch, missing {missing}, unused {unused}")
    return order


def link_doi(s: str) -> str:
    """Make every doi:10.x/... clickable while keeping the printed AMA form."""
    return re.sub(r"doi:(10\.[^\s,]+)",
                  lambda m: f"\\href{{https://doi.org/{m.group(1).replace(chr(92) + '_', '_')}}}{{doi:{m.group(1)}}}", s)


def references_tex(order: list[str]) -> str:
    items = "\n".join(f"\\bibitem{{{k}}} {link_doi(REFS[k])}" for k in order)
    return f"\\begin{{thebibliography}}{{{len(order)}}}\n{items}\n\\end{{thebibliography}}\n"


def compress(nums: list[int]) -> str:
    nums, out, i = sorted(nums), [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"{nums[i]}-{nums[j]}" if j - i >= 2 else ",".join(map(str, nums[i:j + 1])))
        i = j + 1
    return ",".join(out)


def env_body(tex: str, env: str) -> str:
    m = re.search(rf"\\begin\{{{env}\}}(.*?)\\end\{{{env}\}}", tex, re.S)
    return m.group(1).strip() if m else ""


def docx_tex(tex: str, order: list[str]) -> str:
    """Plain LaTeX that pandoc can read, with numbered superscript citations and a numbered reference list."""
    num = {k: i + 1 for i, k in enumerate(order)}
    title = re.search(r"\\title\{(.*?)\}\n\n", tex, re.S).group(1)
    names = [re.sub(r"\\corref\{[^}]*\}", "*", a)
             for a in re.findall(r"\\author\[[^]]*\]\{((?:[^{}]|\{[^{}]*\})*)\}", tex)]
    aff = re.search(r"organization=\{([^}]*)\}", tex).group(1)
    city = re.search(r"city=\{([^}]*)\}", tex).group(1)
    pc = re.search(r"postcode=\{([^}]*)\}", tex).group(1)
    ctry = re.search(r"country=\{([^}]*)\}", tex).group(1)
    emails = re.findall(r"\\ead\{([^}]*)\}", tex)
    cort = re.search(r"\\cortext\[[^]]*\]\{(.*?)\}\n", tex, re.S).group(1)
    authors = (", ".join(f"{n}\\textsuperscript{{a}}" for n in names) + "\n\n\\textsuperscript{a}"
               + f"{aff}, {city} {pc}, {ctry}\n\n* {cort} E-mail addresses " + ", ".join(emails) + ".")
    abstract = env_body(tex, "abstract").replace("\\noindent", "")
    keywords = re.sub(r"\s*\\sep\s*", "; ", env_body(tex, "keyword"))
    body = tex.split("\\end{frontmatter}", 1)[1].split("\\end{document}")[0]
    body = re.sub(r"\\(linenumbers|RaggedRight)\b", "", body)
    body = re.sub(r"\\setlength\{\\parindent\}\{[^}]*\}", "", body)
    body = re.sub(r"\\cite\{([^}]*)\}",
                  lambda m: "\\textsuperscript{" + compress([num[k.strip()] for k in m.group(1).split(",")]) + "}",
                  body)
    # resolve cross-references and number captions, which pandoc leaves as [label] in Word
    labels, counts = {}, {"tab": 0, "fig": 0, "eq": 0, "sec": 0, "alg": 0}
    for kind, name in re.findall(r"\\label\{(tab|fig|eq|sec|alg):([^}]*)\}", tex):
        counts[kind] += 1
        labels[f"{kind}:{name}"] = str(counts[kind])
    sections = [m.start() for m in re.finditer(r"\\section\{", body)]
    for name in re.findall(r"\\label\{(sec:[^}]*)\}", body):
        pos = body.index(f"\\label{{{name}}}")
        labels[name] = str(sum(1 for p in sections if p < pos))
    body = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels[m.group(1)], body)
    body = re.sub(r"\\begin\{algorithm\}.*?\\end\{algorithm\}", algorithm_text, body, flags=re.S)
    na = iter(range(1, 100))
    body = re.sub(r"ALGCAP ", lambda m: f"Algorithm {next(na)}. ", body)
    png6 = os.path.join(OUT, "figures", "Figure6.png")
    body = re.sub(r"\\figorplaceholder\{([^}]*)\}\{[^}]*\}",
                  lambda m: ("\\includegraphics{" + m.group(1) + "}") if os.path.exists(png6)
                  else "[CHECK: Figure 6 is created by notebook 19 on Kaggle]", body)
    body = re.sub(r"\\label\{[^}]*\}", "", body)
    n = {"table": 0, "figure": 0}

    def cap(m):
        env = m.group(1)
        n[env] += 1
        return m.group(0) + ("Table " if env == "table" else "Fig. ") + f"{n[env]}. "
    body = re.sub(r"\\begin\{(table|figure)\}(?:\[[^]]*\])?(?:(?!\\caption\{).)*\\caption\{", cap, body, flags=re.S)
    body = re.sub(r"\\begin\{(table|figure)\}\[[^]]*\]", r"\\begin{\1}", body)
    def plain_cols(m):
        spec = re.sub(r"@\{(?:[^{}]|\{[^{}]*\})*\}", "", m.group(1))
        return "\\begin{tabular}{" + re.sub(r"p\{[^}]*\}", "l", spec) + "}"
    body = re.sub(r"\\begin\{tabular\*\}\{\\linewidth\}\{((?:[^{}]|\{(?:[^{}]|\{[^{}]*\})*\})*)\}",
                  plain_cols, body)
    body = body.replace("\\end{tabular*}", "\\end{tabular}")
    body = re.sub(r"Figure(\d)\.pdf", r"Figure\1.png", body)
    refs = "\\section*{References}\n\\begin{enumerate}\n" + "\n".join(f"\\item {link_doi(REFS[k])}" for k in order) + \
           "\n\\end{enumerate}\n"
    body = body.replace("\\input{references.tex}", refs)
    return ("\\documentclass{article}\n\\usepackage{amsmath,graphicx,booktabs}\n\\graphicspath{{figures/}}\n"
            f"\\title{{{title}}}\n\\begin{{document}}\n\\maketitle\n"
            f"{authors}\n\n"
            f"\\section*{{Abstract}}\n{abstract}\n\n\\textbf{{Keywords.}} {keywords}\n\n{body}\n\\end{{document}}\n")


def format_docx(path: str) -> None:
    """Journal Word layout. Times New Roman 12 pt, left aligned, double spacing, continuous line numbers."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    d = docx.Document(path)
    for st in d.styles:
        if st.type == 1 and st.font is not None:
            st.font.name = "Times New Roman"
            rpr = st.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.append(fonts)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(a), "Times New Roman")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if fonts.get(qn(a)) is not None:
                    del fonts.attrib[qn(a)]
            st.font.color.rgb = None
            if st.name in ("Normal", "Body Text", "First Paragraph", "Compact", "Abstract"):
                st.font.size = Pt(12)
    for p in d.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if p.style.name not in ("Image Caption", "Table Caption", "Captioned Figure"):
            p.paragraph_format.line_spacing = 2.0
    for sec in d.sections:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sec._sectPr.append(ln)
    d.save(path)


def write_zip(sup: str) -> None:
    """intelligent_oncology_latex.zip for Overleaf. main.tex at the root compiles as is, and supplement/ holds the
    supplement source with its montage beside it, plus the compiled supplementary_material.pdf."""
    import zipfile
    path = os.path.join(OUT, "intelligent_oncology_latex.zip")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in ("main.tex", "references.tex", "elsarticle.cls", "elsarticle-num.bst", "highlights.txt",
                  "supplementary_material.pdf", "graphical_abstract.pdf", "graphical_abstract.png"):
            z.write(os.path.join(OUT, f), f)
        for f in sorted(os.listdir(os.path.join(OUT, "figures"))):
            z.write(os.path.join(OUT, "figures", f), f"figures/{f}")
        montage = "top_cross_split_pairs.png"
        tex = open(os.path.join(sup, "supplement.tex")).read().replace(f"../../results/ncc/{montage}", montage)
        z.writestr("supplement/supplement.tex", tex)
        for f in sorted(os.listdir(os.path.join(sup, "tables"))):
            z.write(os.path.join(sup, "tables", f), f"supplement/tables/{f}")
        z.write(os.path.join(ROOT, "results", "ncc", montage), f"supplement/{montage}")


TITLE_PAGE = r"""\documentclass[11pt,a4paper]{article}
\usepackage[T1]{fontenc}
\usepackage{mathptmx}
\usepackage[margin=2.5cm]{geometry}
\usepackage[hidelinks]{hyperref}
\setlength{\parindent}{0pt}
\setlength{\parskip}{8pt}
\pagestyle{empty}
\begin{document}
\section*{Title page}
\textbf{Title.} TITLE

\textbf{Authors.} AUTHORS

\textsuperscript{a}AFFIL

\textbf{Corresponding author.} CORR, AFFIL. E-mail: CORRMAIL. Telephone: ..............................

\textbf{E-mail addresses.} EMAILS
\end{document}
"""

ETHICS = r"""\documentclass[11pt,a4paper]{article}
\usepackage[T1]{fontenc}
\usepackage{mathptmx}
\usepackage[margin=2.5cm]{geometry}
\setlength{\parindent}{0pt}
\pagestyle{empty}
\begin{document}
\section*{Ethical statement}
\textbf{Title.} TITLE

BODY
\end{document}
"""


def section_text(tex: str, name: str) -> str:
    m = re.search(r"\\section\*\{" + re.escape(name) + r"\}\n(.*?)(?=\n\\section|\n\\input)", tex, re.S)
    return m.group(1).strip()


def _plain(t: str) -> str:
    """LaTeX cell or caption text to plain Word text."""
    t = re.sub(r"\$\\varepsilon\s*=\s*([^$]*)\$", r"ε = \1", t)
    t = t.replace("$\\varepsilon$", "ε").replace("$\\sigma$", "σ")
    t = re.sub(r"\\textsuperscript\{([^}]*)\}", "\x01\\1\x02", t)
    for g, u in (("alpha", "α"), ("sigma", "σ"), ("delta", "δ"), ("mu", "μ"), ("varepsilon", "ε"), ("ln", "ln")):
        t = t.replace("\\" + g, u)
    t = re.sub(r"\{\\scriptsize\s*([^}]*)\}", r"\1", t)
    t = re.sub(r"\$([^$]*)\$", r"\1", t)
    t = t.replace("\\times", "×").replace("\\%", "%").replace("~", " ").replace("\\ ", " ")
    t = re.sub(r"\\(textbf|textit|emph|mathrm)\{([^}]*)\}", r"\2", t)
    t = re.sub(r"\\[a-zA-Z]+\*?(\[[^]]*\])?(\{[^}]*\})?", "", t)
    return re.sub(r"\s+", " ", t.replace("{", "").replace("}", "")).strip()


def _runs(par, text: str, bold: bool = False, size: float | None = None):
    """Write text into a Word paragraph, turning \x01x\x02 markers into superscript runs."""
    from docx.shared import Pt
    for k, piece in enumerate(re.split("\x01|\x02", text)):
        if not piece:
            continue
        run = par.add_run(piece)
        run.bold = bold
        run.font.superscript = k % 2 == 1
        if size:
            run.font.size = Pt(size)


def tables_docx(tex: str, dst: str):
    """Tables 1 to N of main.tex as real Word tables (grid, header row, CI under each mean), with captions and notes."""
    import docx
    from docx.shared import Pt
    out = docx.Document()
    st = out.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10)
    out.add_heading("Tables", level=1)
    for n, m in enumerate(re.finditer(r"\\begin\{table\}.*?\\end\{table\}", tex, re.S), 1):
        t = m.group(0)
        cm = re.search(r"\\caption\{(.*?)\}\s*(\\label|\n)", t, re.S)
        cap = _plain(cm.group(1))
        body = re.search(r"\\toprule(.*?)\\bottomrule", t, re.S).group(1)
        rows = []
        for line in re.split(r"\\\\(?:\[[^]]*\])?", body):
            line = line.replace("\\midrule", "").strip()
            if not line:
                continue
            cells = [_plain(c) for c in line.split("&")]
            if rows and cells[0] == "" and all(c == "" or c.startswith("(") for c in cells[1:]) and any(cells[1:]):
                rows[-1] = [r + (" " + c if c else "") for r, c in zip(rows[-1], cells)]
            else:
                rows.append(cells)
        p = out.add_paragraph()
        p.add_run(f"Table {n}. ").bold = True
        _runs(p, cap)
        ncol = max(len(r) for r in rows)
        tb = out.add_table(rows=len(rows), cols=ncol)
        tb.style = "Table Grid"
        for i, r in enumerate(rows):
            for j in range(ncol):
                cell = tb.cell(i, j)
                _runs(cell.paragraphs[0], r[j] if j < len(r) else "", bold=i == 0)
        note = re.search(r"\\par\\smallskip\{\\scriptsize(.*?)\}\n\\end\{table\}", t, re.S)
        if note:
            _runs(out.add_paragraph(), _plain(note.group(1)), size=9)
        out.add_paragraph("")
    out.save(dst)


def em_package(tex: str):
    """Flat folder for Editorial Manager, which rejects zip files with subfolders."""
    em = os.path.join(OUT, "em_upload")
    os.makedirs(em)
    title = re.sub(r"\s+", " ", re.search(r"\\title\{(.*?)\}\n\n", tex, re.S).group(1))
    names = [re.sub(r"\\corref\{[^}]*\}", "", a) for a in re.findall(r"\\author\[[^]]*\]\{((?:[^{}]|\{[^{}]*\})*)\}", tex)]
    emails = re.findall(r"\\ead\{([^}]*)\}", tex)
    corr = [re.sub(r"\\corref\{[^}]*\}", "", a) for a in re.findall(r"\\author\[[^]]*\]\{((?:[^{}]|\{[^{}]*\})*)\}", tex)
            if "corref" in a][0]
    affil = "School of Computer and Systems Sciences, Jawaharlal Nehru University, New Delhi 110067, India"
    tp = (TITLE_PAGE.replace("TITLE", title).replace("AUTHORS", ", ".join(f"{n}\\textsuperscript{{a}}" for n in names))
          .replace("CORRMAIL", emails[names.index(corr)].replace("_", "\\_")).replace("CORR", corr).replace("AFFIL", affil)
          .replace("EMAILS", ", ".join(f"{e.replace('_', chr(92) + '_')} ({n})" for n, e in zip(names, emails))))
    eth = ETHICS.replace("TITLE", title).replace("BODY", section_text(tex, "Ethic statement"))
    for name, src in (("title_page", tp), ("ethical_statement", eth)):
        open(os.path.join(em, f"{name}.tex"), "w").write(src)
        run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", f"{name}.tex"], em)
        run(["pandoc", f"{name}.tex", "-f", "latex", "-t", "docx", "-o", f"{name}.docx"], em)
        for ext in ("aux", "log", "out", "tex"):
            p = os.path.join(em, f"{name}.{ext}")
            if os.path.exists(p):
                os.remove(p)
    tables_docx(tex, os.path.join(em, "tables_editable.docx"))
    for f in ("manuscript.pdf", "cover_letter.pdf", "cover_letter.docx", "highlights.docx",
              "graphical_abstract.png", "supplementary_material.pdf"):
        shutil.copy(os.path.join(OUT, f), em)
    for f in sorted(os.listdir(os.path.join(OUT, "figures"))):
        if f.startswith("Figure") and f.endswith(".pdf"):
            shutil.copy(os.path.join(OUT, "figures", f), em)
    # LaTeX source as a flat zip (asked for at revision): graphicx finds figures in the same folder
    with zipfile.ZipFile(os.path.join(em, "latex_source_flat.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        for f in ("main.tex", "references.tex", "elsarticle.cls", "elsarticle-num.bst"):
            z.write(os.path.join(OUT, f), f)
        for f in sorted(os.listdir(os.path.join(OUT, "figures"))):
            z.write(os.path.join(OUT, "figures", f), f)


def run(cmd, cwd, env=None):
    r = subprocess.run(cmd, cwd=cwd, env=env, capture_output=True, text=True, errors="replace")
    if r.returncode:
        print(r.stdout[-3000:], r.stderr[-2000:])
        sys.exit(f"failed: {' '.join(cmd)}")


def main():
    tex = open(os.path.join(HERE, "main.tex")).read()
    order = cite_order(tex)
    if os.path.exists(OUT):
        shutil.rmtree(OUT)
    os.makedirs(os.path.join(OUT, "figures"))
    open(os.path.join(HERE, "references.tex"), "w").write(references_tex(order))
    for f in ("main.tex", "references.tex"):
        shutil.copy(os.path.join(HERE, f), OUT)
    for f in ("elsarticle.cls", "elsarticle-num.bst"):
        shutil.copy(os.path.join(HERE, "elsarticle", f), OUT)
    for new, old in FIGS.items():
        for ext in ("pdf", "png"):
            src = os.path.join(ROOT, "figures", f"{old}.{ext}")
            if os.path.exists(src):
                shutil.copy(src, os.path.join(OUT, "figures", f"{new}.{ext}"))
            else:
                print(f"note: {old}.{ext} not yet produced (Kaggle notebook 19), {new} is a placeholder")
    for new, src in TIKZ.items():
        tikz_pdf(src, os.path.join(OUT, "figures"), new)
    tikz_pdf("graphical_abstract", OUT, "graphical_abstract")
    shutil.copy(os.path.join(OUT, "graphical_abstract.pdf"), os.path.join(OUT, "figures", "graphical_abstract.pdf"))
    hl = [re.sub(r"\$\\varepsilon\s*=\s*(\d+)\$", r"ε = \1", h.strip()).replace("$\\varepsilon$", "ε")
          for h in env_body(tex, "highlights").split("\\item") if h.strip()]
    too_long = [h for h in hl if len(h) > 85]
    if too_long or not 3 <= len(hl) <= 5:
        sys.exit(f"highlights must be 3 to 5 items of at most 85 characters, got {too_long}")
    open(os.path.join(OUT, "highlights.txt"), "w").write("Highlights\n\n" + "\n".join(f"• {h}" for h in hl) + "\n")

    env = dict(os.environ, TEXINPUTS=os.path.join(HERE, "elsarticle") + "//:")
    run(["latexmk", "-pdf", "-interaction=nonstopmode", "-halt-on-error", "main.tex"], OUT, env)
    shutil.copy(os.path.join(OUT, "main.pdf"), os.path.join(OUT, "manuscript.pdf"))

    open(os.path.join(OUT, "_docx.tex"), "w").write(docx_tex(tex, order))
    run(["pandoc", "_docx.tex", "-f", "latex", "-t", "docx", "-o", "manuscript.docx", "--resource-path=.:figures"],
        OUT)
    format_docx(os.path.join(OUT, "manuscript.docx"))

    run([sys.executable, os.path.join(ROOT, "scripts", "make_supplement.py")], ROOT)
    sup = os.path.join(HERE, "supplement")
    run(["latexmk", "-pdf", "-interaction=nonstopmode", "-halt-on-error", "supplement.tex"], sup)
    shutil.copy(os.path.join(sup, "supplement.pdf"), os.path.join(OUT, "supplementary_material.pdf"))

    run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "-output-directory", OUT,
         os.path.join(HERE, "cover_letter.tex")], HERE)
    run(["pandoc", os.path.join(HERE, "cover_letter.tex"), "-f", "latex", "-t", "docx", "-o", "cover_letter.docx"], OUT)
    run(["pandoc", "highlights.txt", "-f", "markdown", "-t", "docx", "-o", "highlights.docx"], OUT)
    for f in os.listdir(OUT):
        if f.startswith("cover_letter.") and f.split(".")[-1] in ("aux", "log", "out"):
            os.remove(os.path.join(OUT, f))
    for f in os.listdir(OUT):
        if f.startswith("main.") and f not in ("main.tex",) or f.startswith("_docx"):
            os.remove(os.path.join(OUT, f))
    write_zip(sup)
    em_package(tex)
    print(f"{len(order)} references, {len(hl)} highlights")
    print("\n".join(sorted(os.path.relpath(os.path.join(dp, f), OUT) for dp, _, fs in os.walk(OUT) for f in fs)))


if __name__ == "__main__":
    main()
